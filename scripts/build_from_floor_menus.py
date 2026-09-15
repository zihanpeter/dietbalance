"""根据四份楼层菜单重建 ``data/dishes.json``。

流程：
1. 从一层/二层/三层/主食早餐菜单提取全部菜名
2. 用现有 dishes.json 补宏量素占比与菜品特点
3. 用「菜品热量估算表」补一份总热量、出品重量、类别、每100g热量
4. 输出以当前菜单为准的菜品库（菜单上有的菜都会收录）
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl
import xlrd

ROOT = Path(__file__).resolve().parent.parent
OUT_JSON = ROOT / "data" / "dishes.json"
OLD_DISHES = ROOT / "data" / "dishes.json"
KCAL_XLSX = Path(
    r"C:\Users\luzih\Documents\xwechat_files\wxid_kzabuv4uymt322_792e\msg\file\2026-08\菜品热量估算表.xlsx"
)
BASE = Path(
    r"C:\Users\luzih\Documents\xwechat_files\wxid_kzabuv4uymt322_792e"
    r"\msg\attach\27f1dfd493539e7170717ea621624c08\2026-09\Rec\e90021652ac2f8ec\F"
)

sys.stdout.reconfigure(encoding="utf-8")

SKIP_EXACT = {
    "",
    "食品名称",
    "菜品",
    "热菜",
    "盖饭",
    "一层午餐菜单",
    "一层晚餐菜单",
    "二层菜谱",
    "三层菜单",
    "早餐菜单",
    "星期一",
    "星期二",
    "星期三",
    "星期四",
    "星期五",
    "周一",
    "周二",
    "周三",
    "周四",
    "周五",
}

# 菜单常见异体字 / 错别字 -> 规范名片段
ALIAS_REPLACEMENTS = (
    ("黒椒", "黑椒"),
    ("千页", "千叶"),
    ("西胡芦", "西葫芦"),
    ("肉未", "肉末"),
    ("祘茸", "蒜蓉"),
    ("蒜茸", "蒜蓉"),
    ("蕃茄", "番茄"),
    ("牛楠蛊", "牛腩盅"),
    ("虫草牛楠蛊", "虫草牛腩盅"),
    ("密制", "蜜制"),
    ("咖哩", "咖喱"),
    ("氽", "汆"),
    ("元白菜", "圆白菜"),
    ("栆卷", "枣卷"),
    ("西蓝花", "西兰花"),
)


def normalize(name: object) -> str:
    text = str(name or "").strip()
    text = re.sub(r"\s+", "", text)
    for a, b in ALIAS_REPLACEMENTS:
        text = text.replace(a, b)
    return text


def is_skip(cell: object) -> bool:
    if cell is None:
        return True
    text = str(cell).strip()
    if not text or text in SKIP_EXACT:
        return True
    if text.startswith("档口"):
        return True
    if "菜单" in text or "菜谱" in text:
        return True
    return False


def extract_xls(path: Path, floor: str) -> dict[str, set[str]]:
    wb = xlrd.open_workbook(path)
    found: dict[str, set[str]] = {}
    for sheet_name in wb.sheet_names():
        ws = wb.sheet_by_name(sheet_name)
        tag = f"{floor}/{sheet_name}"
        for r in range(ws.nrows):
            for c in range(ws.ncols):
                cell = ws.cell_value(r, c)
                if is_skip(cell):
                    continue
                name = normalize(cell)
                if not name or name in SKIP_EXACT:
                    continue
                found.setdefault(name, set()).add(tag)
    return found


def extract_xlsx(path: Path, floor: str) -> dict[str, set[str]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    found: dict[str, set[str]] = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        tag = f"{floor}/{sheet_name}"
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if is_skip(cell):
                    continue
                name = normalize(cell)
                if not name or name in SKIP_EXACT:
                    continue
                found.setdefault(name, set()).add(tag)
    return found


def load_old_dishes() -> dict[str, dict]:
    if not OLD_DISHES.exists():
        return {}
    raw = json.loads(OLD_DISHES.read_text(encoding="utf-8"))
    return {normalize(d["name"]): d for d in raw if d.get("name")}


def load_kcal_table(path: Path) -> dict[str, dict]:
    if not path.exists():
        print(f"kcal table missing: {path}")
        return {}
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["全部菜品热量"]
    mapping: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        name = normalize(row[0])
        mapping[name] = {
            "category": str(row[1]).strip() if row[1] else None,
            "portion_g": _to_int(row[6]),
            "total_kcal": _to_int(row[9]),
            "kcal": _to_int(row[10]),  # 每100g
        }
    print(f"Loaded kcal table: {len(mapping)} dishes")
    return mapping


def _to_int(value: object) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return None


def _to_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def resolve_enrichment(
    name: str,
    old: dict[str, dict],
    kcal_map: dict[str, dict],
) -> tuple[dict | None, dict | None, str]:
    """返回 (old_hit, kcal_hit, match_note)。"""
    if name in old and name in kcal_map:
        return old[name], kcal_map[name], "exact"
    if name in old:
        return old[name], None, "old_only"
    if name in kcal_map:
        return None, kcal_map[name], "kcal_only"

    # 唯一子串匹配
    old_cands = [k for k in old if name in k or k in name]
    kcal_cands = [k for k in kcal_map if name in k or k in name]
    old_hit = old[old_cands[0]] if len(old_cands) == 1 else None
    kcal_hit = kcal_map[kcal_cands[0]] if len(kcal_cands) == 1 else None
    if old_hit or kcal_hit:
        note = "fuzzy"
        if old_cands:
            note += f"/old:{old_cands[0]}"
        if kcal_cands:
            note += f"/kcal:{kcal_cands[0]}"
        return old_hit, kcal_hit, note
    return None, None, "none"


def main() -> None:
    menu: dict[str, set[str]] = {}
    for part in (
        extract_xls(BASE / "0" / "一层菜单.xls", "一层"),
        extract_xlsx(BASE / "1" / "二层菜谱.xlsx", "二层"),
        extract_xlsx(BASE / "2" / "三层菜单.xlsx", "三层"),
        extract_xlsx(BASE / "3" / "主食早餐.xlsx", "主食早餐"),
    ):
        for name, tags in part.items():
            menu.setdefault(name, set()).update(tags)

    old = load_old_dishes()
    kcal_map = load_kcal_table(KCAL_XLSX)

    dishes: list[dict] = []
    stats = {"exact": 0, "old_only": 0, "kcal_only": 0, "fuzzy": 0, "none": 0}

    for name in sorted(menu.keys()):
        old_hit, kcal_hit, note = resolve_enrichment(name, old, kcal_map)
        bucket = note.split("/")[0]
        stats[bucket] = stats.get(bucket, 0) + 1

        floors = sorted({t.split("/")[0] for t in menu[name]})

        # 展示名优先用旧库里的规范写法，否则用菜单归一化名
        display_name = str(old_hit["name"]).strip() if old_hit and old_hit.get("name") else name

        kcal = None
        total_kcal = None
        portion_g = None
        category = None
        carb_pct = protein_pct = fat_pct = 0.0
        features = ""

        if old_hit:
            kcal = _to_int(old_hit.get("kcal"))
            carb_pct = float(old_hit.get("carb_pct") or 0)
            protein_pct = float(old_hit.get("protein_pct") or 0)
            fat_pct = float(old_hit.get("fat_pct") or 0)
            features = str(old_hit.get("features") or "").strip()
            if old_hit.get("total_kcal") is not None:
                total_kcal = _to_int(old_hit.get("total_kcal"))
            if old_hit.get("portion_g") is not None:
                portion_g = _to_int(old_hit.get("portion_g"))
            if old_hit.get("category"):
                category = str(old_hit["category"]).strip()

        if kcal_hit:
            if kcal_hit.get("kcal") is not None:
                kcal = kcal_hit["kcal"]
            if kcal_hit.get("total_kcal") is not None:
                total_kcal = kcal_hit["total_kcal"]
            if kcal_hit.get("portion_g") is not None:
                portion_g = kcal_hit["portion_g"]
            if kcal_hit.get("category"):
                category = kcal_hit["category"]

        # 每100g 热量兜底：用一份总热量 / 出品重量估算
        if kcal is None and total_kcal and portion_g:
            kcal = int(round(total_kcal * 100 / portion_g))

        dishes.append(
            {
                "name": display_name,
                "kcal": kcal or 0,
                "carb_pct": round(carb_pct, 2),
                "protein_pct": round(protein_pct, 2),
                "fat_pct": round(fat_pct, 2),
                "features": features,
                "total_kcal": total_kcal,
                "portion_g": portion_g,
                "category": category,
                "floors": floors,
            }
        )

    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(dishes, ensure_ascii=False, indent=2), encoding="utf-8")

    with_total = sum(1 for d in dishes if d.get("total_kcal"))
    with_macros = sum(1 for d in dishes if d.get("carb_pct") or d.get("protein_pct") or d.get("fat_pct"))
    print(f"Wrote {len(dishes)} dishes -> {OUT_JSON}")
    print(f"with total_kcal={with_total}, with macros={with_macros}")
    print(f"match stats={stats}")
    sample = next((d for d in dishes if d["name"] == "宫保鸡丁"), None)
    print(f"sample 宫保鸡丁: {sample}")


if __name__ == "__main__":
    main()
